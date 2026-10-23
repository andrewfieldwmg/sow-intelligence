"""Builds the fictional demo corpus: 48 SOWs (38 Adobe implementations, 10 Applied AI engagements) (Word and PDF) for invented clients, plus the alias
map, the delivery register and a golden eval set whose answers are known by construction.

Every company, person and number here is made up. Clause wording is illustrative, written to look
like UK practice, and is not legal advice.

Sector clause libraries come in strong and weak versions. Strong versions mostly land in deals the
register marks green and on-margin; weak ones mostly in amber or red deals. That is what lets the
app answer "what works" from outcomes rather than opinion.
"""
import csv
import json
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(1, str(Path(__file__).resolve().parent.parent))
import boilerplate as bp                                   # noqa: E402
from ingest import read_pages                              # noqa: E402

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "demo")
rng = random.Random(42)

# (name, short name used in prose or None, sector, size). All invented.
CLIENTS = [
    ("Brackwater Building Society", "Brackwater", "financial services", "enterprise"),
    ("Quillon Insurance plc", "Quillon", "financial services", "enterprise"),
    ("Belmont Mutual", None, "financial services", "enterprise"),
    ("Harrowgate Bank plc", "Harrowgate", "financial services", "enterprise"),
    ("Kestrel Life Assurance", "Kestrel", "financial services", "enterprise"),
    ("Northdene Credit Union", "Northdene", "financial services", "smb"),
    ("Selby Wealth Management", None, "financial services", "mid-market"),
    ("Tolland Payments Ltd", "Tolland", "financial services", "mid-market"),
    ("Arden Capital Markets", None, "financial services", "enterprise"),
    ("Cawston Pensions Trust", "Cawston", "financial services", "mid-market"),
    ("Halden Retail Group plc", "Halden", "retail", "enterprise"),
    ("Pellow Home Stores", "Pellow", "retail", "mid-market"),
    ("Vexley Fashion Ltd", "Vexley", "retail", "mid-market"),
    ("Orrin Grocers plc", "Orrin", "retail", "enterprise"),
    ("Larkfield Outdoor", "Larkfield", "retail", "mid-market"),
    ("Thistlebrook Pharmacy", "Thistlebrook", "retail", "enterprise"),
    ("Calderwick Electricals", "Calderwick", "retail", "enterprise"),
    ("Penhale Books", "Penhale", "retail", "smb"),
    ("Rushmere Furniture", "Rushmere", "retail", "mid-market"),
    ("Zennor Hotels Ltd", "Zennor", "hospitality", "mid-market"),
    ("Marlowe Leisure Parks", "Marlowe", "hospitality", "enterprise"),
    ("Glenvar Resorts", "Glenvar", "hospitality", "enterprise"),
    ("Saltmarsh Inns", None, "hospitality", "mid-market"),
    ("Brightwater Cruises", None, "hospitality", "enterprise"),
    ("Corran Holidays", "Corran", "hospitality", "mid-market"),
    ("Ashby Hotel Group", None, "hospitality", "enterprise"),
    ("Lowther Spa and Golf", "Lowther", "hospitality", "smb"),
    ("Corrin Motors Ltd", "Corrin", "automotive", "enterprise"),
    ("Felsham Pharma AG", "Felsham", "pharma", "enterprise"),
    ("Ostrava Components plc", "Ostrava", "industrial", "enterprise"),
    ("Ravensby Telecom", "Ravensby", "telecoms", "enterprise"),
    ("Stirlow Foods Ltd", "Stirlow", "consumer goods", "mid-market"),
    ("Tamsin Optics GmbH", "Tamsin", "healthcare", "mid-market"),
    ("Ulverton Energy plc", "Ulverton", "energy", "enterprise"),
    ("Wendover Rail", "Wendover", "transport", "enterprise"),
    ("Yardley Learning Trust", "Yardley", "education", "smb"),
    ("Ashcombe Distribution", "Ashcombe", "logistics", "mid-market"),
    ("Felbridge Water", "Felbridge", "utilities", "enterprise"),
    # Applied AI engagements: forward deployed engineers building production AI with the client's team.
    ("Ferncliffe Bank plc", "Ferncliffe", "financial services", "enterprise"),
    ("Oakhurst Insurance Group", "Oakhurst", "financial services", "enterprise"),
    ("Mallory Lending Ltd", "Mallory", "financial services", "mid-market"),
    ("Draycott Department Stores", "Draycott", "retail", "enterprise"),
    ("Kinver Sportswear", "Kinver", "retail", "mid-market"),
    ("Harbourlight Hotels", "Harbourlight", "hospitality", "enterprise"),
    ("Eskdale Travel Group", "Eskdale", "hospitality", "mid-market"),
    ("Penrose Health Partners", "Penrose", "healthcare", "enterprise"),
    ("Vantor Logistik GmbH", "Vantor", "logistics", "enterprise"),
    ("Ribblesdale Housing Association", "Ribblesdale", "public sector", "mid-market"),
]
AI_FROM = 38                                              # clients from this index are Applied AI engagements
EU_CLIENTS = {"Felsham Pharma AG", "Tamsin Optics GmbH", "Arden Capital Markets", "Tolland Payments Ltd",
              "Vantor Logistik GmbH"}
CONTACTS = ["Priya Natarajan", "Tom Ellery", "Sian Morgan", "Lukas Brandt", "Ana Ferreira", "Dev Malhotra",
            "Grace Okafor", "Owen Pryce"]
# Delivery outcome per client index: mostly green, some amber, a few red.
# Pre-sales "investment" engagements: no fee, funded by the Supplier to prove value and convert.
INVESTMENT = {2, 9, 17, 24, 31, 37, 38, 41, 42, 45, 46}
NOT_CONVERTED = {9, 24, 46}           # met their success criteria but the client didn't fund the next phase
HEALTH = {i: "red" for i in (4, 12, 15, 22, 30, 35, 43)} | {i: "amber" for i in (1, 7, 10, 18, 20, 25, 28, 33, 41, 45)}

PRODUCTS = {
    "AEM Sites": {
        "deliverables": [
            ("component library of {n} reusable AEM Sites components built to the agreed design system", {"architect": 60, "developer": 420, "qa": 80}),
            ("migration of {n} pages from the legacy CMS to AEM Sites as a Cloud Service, including redirects", {"developer": 260, "qa": 60}),
            ("headless delivery of content to the existing React front end using AEM content fragments and GraphQL", {"architect": 40, "developer": 220}),
            ("multi-site manager set-up for {n} markets with language copies and a translation workflow", {"architect": 30, "developer": 160, "qa": 40}),
        ],
        "assumptions": ["The Client supplies final designs in Figma before the build sprint for each template starts.",
                        "Content owners in each market are available to sign off migrated pages within five working days.",
                        "The Client's existing CDN and DNS changes are made by the Client's infrastructure team."],
        "out_of_scope": ["Changes to the React front end beyond consuming the agreed GraphQL endpoints.",
                         "Translation of content, which the Client's agency provides."],
    },
    "AEM Assets": {
        "deliverables": [
            ("metadata schema and folder taxonomy for {n} asset types, agreed with brand and legal", {"architect": 30, "consultant": 60}),
            ("automated asset ingestion and approval workflow from the creative agency into AEM Assets", {"architect": 20, "developer": 140, "qa": 30}),
            ("Brand Portal set-up so that {n} partner agencies can download approved assets", {"developer": 60, "consultant": 20}),
        ],
        "assumptions": ["The Client migrates historic assets using its own team; the Supplier provides the migration scripts.",
                        "Rights and expiry data exists for every asset to be migrated."],
        "out_of_scope": ["Migration of video assets over 2GB.", "Integration with third-party digital asset managers."],
    },
    "Target": {
        "deliverables": [
            ("{n} personalisation activities in Adobe Target, from hypothesis to live test and readout", {"consultant": 120, "developer": 80}),
            ("audience strategy linking Target to the Client's customer data platform for real-time segments", {"architect": 40, "consultant": 60}),
            ("testing playbook and prioritisation framework handed over to the Client's optimisation team", {"consultant": 40}),
        ],
        "assumptions": ["The Client's data layer is already implemented on all in-scope pages.",
                        "Traffic volumes support statistically valid tests within four weeks per activity."],
        "out_of_scope": ["Building or changing the customer data platform.", "Creative production for test variants."],
    },
    "Analytics": {
        "deliverables": [
            ("solution design reference and data layer specification for {n} key journeys", {"architect": 50, "consultant": 70}),
            ("Adobe Analytics implementation through Adobe Launch tags, with QA across all in-scope journeys", {"developer": 160, "qa": 60}),
            ("{n} Analysis Workspace dashboards for marketing, ecommerce and leadership", {"consultant": 90}),
        ],
        "assumptions": ["The Client's developers implement the data layer to the specification.",
                        "One reporting suite per market is sufficient."],
        "out_of_scope": ["Historic data migration from the previous analytics tool."],
    },
    "Journey Optimizer": {
        "deliverables": [
            ("{n} lifecycle journeys in Adobe Journey Optimizer, including welcome, browse abandonment and renewal", {"architect": 40, "developer": 200, "consultant": 60}),
            ("email and push templates built to the brand guidelines, with dynamic content blocks", {"developer": 120, "qa": 40}),
            ("consent and frequency rules aligned with the Client's preference centre", {"architect": 20, "consultant": 30}),
        ],
        "assumptions": ["Customer profile and event data is available in Adobe Experience Platform before journey build starts.",
                        "The Client's ESP domains are warmed before go-live."],
        "out_of_scope": ["SMS channel set-up.", "Decommissioning of the previous email platform."],
    },
    "Applied AI": {
        "deliverables": [
            ("production AI agent for {n} customer-service intents, integrated with the Client's CRM and knowledge base through approved tool calls", {"architect": 40, "fde": 360, "qa": 60}),
            ("evaluation harness and a golden set of at least 200 labelled cases, run on every change and before any model version change", {"fde": 120, "consultant": 40}),
            ("guardrails, escalation to human agents and audit logging for every model-assisted decision", {"architect": 30, "fde": 140, "qa": 40}),
            ("document extraction pipeline for {n} document types, with confidence scores and human review below the agreed threshold", {"fde": 240, "qa": 50}),
            ("enablement programme pairing Supplier engineers with {n} Client engineers, with a starter repository and a playbook for the next use case", {"fde": 100, "consultant": 40}),
        ],
        # Strong deals secure data access and a product owner up front; weak deals don't.
        "assumptions": ["The Client provides read access to a representative, de-identified sample of production data in a sandbox by the end of week two.",
                        "The Client names a product owner with authority to make scope decisions, available at least two days a week.",
                        "Model usage is billed to the Client directly under its own model provider agreement and is not included in the Charges.",
                        "Subject-matter experts label the golden set of evaluation cases within the first three weeks.",
                        "The Client's security review of the solution architecture completes within ten working days of submission."],
        "out_of_scope": ["Fine-tuning or training of models.", "Ongoing model usage costs.",
                         "Organisation-wide change management beyond the pilot team."],
    },
}
ADOBE = [k for k in PRODUCTS if k != "Applied AI"]

# ---- Sector and general clause libraries: strong versions first, weak versions second ----------
FSI_REGULATORY_STRONG = [
    "Regulatory assumption: the Client is responsible for assessing whether these services are a material outsourcing under the FCA's SYSC 8 rules and PRA Supervisory Statement SS2/21; the Supplier will provide the information the Client reasonably needs for that assessment within ten working days of request.",
    "Regulatory assumption: the Client, its auditors and its regulators (including the FCA and PRA) have the access, audit and information rights set out in the regulatory schedule of the master services agreement, and the Supplier will flow those rights down to any subcontractor.",
    "Regulatory assumption: the Client remains responsible for mapping its important business services and impact tolerances under the operational resilience rules; the Supplier will support scenario testing of the in-scope platform on reasonable notice, at the rates in this SOW.",
    "Regulatory assumption: customer data processed under this SOW is stored and processed only in the UK or EEA, and any change of hosting region needs the Client's prior written approval.",
    "Regulatory assumption: the Client is responsible for testing customer communications against its Consumer Duty obligations; the Supplier builds journeys and content to the rules and approvals the Client provides.",
    "Regulatory assumption: the Supplier maintains a documented exit plan for these services, reviewed annually, so the Client can transfer them to itself or another provider in an orderly way.",
]
DORA_STRONG = "Regulatory assumption: as the Client is an EU financial entity, the Supplier will provide the information the Client needs for its register of ICT third-party arrangements under DORA and will report ICT incidents affecting the services within four hours of becoming aware of them."
FSI_REGULATORY_WEAK = [
    "Regulatory assumption: the Supplier will ensure that the solution complies with all applicable laws and regulations.",
    "Regulatory assumption: the Supplier will support the Client with any regulatory requests as required.",
]

ACCEPTANCE_STRONG = [
    "Acceptance: each deliverable is tested against the acceptance criteria and test scripts agreed in writing before its build sprint starts. The Client has ten working days from delivery to run the tests. A deliverable is accepted when no severity 1 or 2 defects remain open; severity 3 and 4 defects are logged and fixed in the next release and do not block acceptance. If the Client neither accepts nor raises a severity 1 or 2 defect within the test window, the deliverable is deemed accepted.",
    "Acceptance: acceptance criteria are objective and measurable and are recorded in the acceptance schedule. The Client may reject a deliverable only by giving written reasons that refer to a specific criterion. The Supplier will fix and resubmit within five working days. After two rejection cycles for the same deliverable, either party may refer the matter to the steering committee, and use of the deliverable in production counts as acceptance.",
]
ACCEPTANCE_WEAK = [
    "Acceptance: each deliverable will be accepted when it is to the Client's reasonable satisfaction.",
    "Acceptance: the Client will confirm acceptance once it is happy that the deliverable meets its requirements.",
]

PAYMENT_FIXED_STRONG = [
    ("15% on mobilisation, 25% on design sign-off, 30% on completion of user acceptance testing, 20% on go-live, and 10% retention released 30 days after go-live once hypercare exits", "15/25/30/20/10 with retention"),
    ("20% on mobilisation, 30% on design sign-off, 30% on UAT completion and 20% on go-live; milestones are deemed achieved if the Client has not raised a blocking defect within ten working days", "20/30/30/20 with deemed achievement"),
]
PAYMENT_FIXED_WEAK = [
    ("10% on signature and 90% on final acceptance of all deliverables", "10/90 back-loaded"),
    ("100% on final acceptance of the complete solution", "100% on completion"),
]
RETAIL_PEAK = "Assumption: no production releases take place during the Client's peak trading freeze from 1 November to 15 January; milestones that fall in the freeze move to the first working week after it, and the payment schedule moves with them."

SUBCONTRACTING_STRONG = [
    "Subcontracting: the Supplier may subcontract parts of the services only with the Client's prior written consent, which is not to be unreasonably withheld. The approved subcontractors are listed in Appendix B. The Supplier remains responsible for the acts and omissions of its subcontractors as if they were its own, and will flow down obligations of confidentiality, data protection and information security no less protective than those in this SOW.",
    "Subcontracting: where a subcontractor will handle payment card data, for example in booking or point-of-sale integrations, the Supplier will ensure the subcontractor is PCI DSS compliant and will provide its attestation of compliance on request. No subcontractor may process guest personal data outside the UK or EEA without the Client's written approval.",
    "Subcontracting: to cover the Client's peak season from May to September, the Supplier may add named subcontractors from Appendix B on ten working days' notice; replacements must have equivalent skills and experience, and the Supplier bears any handover cost.",
]
SUBCONTRACTING_WEAK = [
    "Subcontracting: the Supplier may subcontract any part of the services at its discretion.",
]

STAFFING_AI = "The Supplier will provide a forward deployed team of an engagement manager, an applied AI architect and {dev} forward deployed engineers, working on site with the Client's engineers at least three days a week."
STAFFING = "The Supplier will provide a delivery team of an engagement manager, a solution architect and {dev} developers, with QA support as needed."
GOVERNANCE = "Weekly status meetings review progress, RAID and budget, and a monthly steering committee with the Client sponsor approves changes to priorities."
CHANGE = "Any change to scope, timeline or fees is handled through a written change request, assessed by the Supplier within five working days and approved by both parties before work starts."


def _letters(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


# Letter pairs Postgres drops as English stop words; an alias like "Client AS" would be unsearchable.
STOP = {"AM", "AN", "AS", "AT", "BE", "BY", "DO", "HE", "IF", "IN", "IS", "IT", "ME", "MY", "NO", "OF", "ON",
        "OR", "SO", "TO", "UP", "US", "WE"}
ALIASES = [a for a in (_letters(k) for k in range(80)) if a not in STOP]


def build():
    (OUT / "sows").mkdir(parents=True, exist_ok=True)
    aliases, register, golden, facts = [], [], [], []
    for i, (name, short, sector, size) in enumerate(CLIENTS):
        alias = f"Client {ALIASES[i]}"
        aliases.append((name, alias))
        if short:
            aliases.append((short, alias))
        health = HEALTH.get(i, "green")
        strong = health == "green" or (health == "amber" and i % 2 == 0)
        ai = i >= AI_FROM
        product = "Applied AI" if ai else ADOBE[i % len(ADOBE)]
        p = PRODUCTS[product]
        year = 2022 + i % 4
        start = date(year, 1 + (i * 5) % 12, 1)
        end = start + timedelta(days=rng.choice([120, 180, 240, 300]))
        pricing = ("investment" if i in INVESTMENT else
                   "fixed_price" if (sector == "retail" or rng.random() < 0.6) else "t_and_m")
        dels = rng.sample(p["deliverables"], k=min(len(p["deliverables"]), rng.choice([2, 3])))
        effort_total, del_lines = {}, []
        for text, eff in dels:
            n = rng.choice([4, 6, 8, 12, 20, 40])
            scale = rng.choice([0.85, 1.0, 1.15])
            hrs = {r: round(h * scale) for r, h in eff.items()}
            for r, h in hrs.items():
                effort_total[r] = effort_total.get(r, 0) + h
            role_txt = ", ".join(f"{r.upper() if len(r) <= 3 else r.title()} {h}h" for r, h in hrs.items())
            del_lines.append(f"The Supplier will deliver: {text.format(n=n)} (estimated effort: {role_txt}).")
        pm = round(sum(effort_total.values()) * rng.choice([0.10, 0.12, 0.15]))
        effort_total["pm"] = pm
        rate = rng.choice([1050, 1150, 1250])
        value = round(sum(effort_total.values()) / 8 * rate, -3)
        currency = "EUR" if name in EU_CLIENTS else "GBP"
        contact = CONTACTS[i % len(CONTACTS)]
        handle = (short or name.split()[0]).lower()

        # Sector-specific material, strong or weak depending on how the deal went.
        extra_assumptions, sector_clauses, tags = [], [], {}
        if sector == "financial services":
            if strong:
                extra_assumptions += rng.sample(FSI_REGULATORY_STRONG, k=4)
                if name in EU_CLIENTS:
                    extra_assumptions.append(DORA_STRONG)
            else:
                extra_assumptions.append(rng.choice(FSI_REGULATORY_WEAK))
            tags["regulatory"] = "strong" if strong else "weak"
        if sector == "retail":
            if strong:
                extra_assumptions.append(RETAIL_PEAK)
            tags["peak_freeze"] = strong
        if sector == "hospitality" or (sector in ("transport", "telecoms") and rng.random() < 0.5):
            subs = rng.sample(SUBCONTRACTING_STRONG, k=2) if strong else SUBCONTRACTING_WEAK[:]
            sector_clauses += subs
            tags["subcontracting"] = "strong" if strong else "weak"
        tags["acceptance"] = "strong" if strong else "weak"
        if pricing == "fixed_price":
            sched, label = rng.choice(PAYMENT_FIXED_STRONG if strong else PAYMENT_FIXED_WEAK)
            payment = (f"Payment. The fixed price is invoiced against the Milestones in Schedule 2. "
                       f"Invoices are payable within {rng.choice([30, 45])} days of receipt.")
            tags["payment_schedule"] = label
        elif pricing == "investment":
            sched = None
            payment = rng.choice(bp.INVESTMENT_STRONG if strong else bp.INVESTMENT_WEAK)
            tags["investment"] = "strong" if strong else "weak"
        else:
            sched = None
            payment = (f"Payment. Time and materials are invoiced monthly in arrears against approved timesheets; "
                       f"invoices are payable within {rng.choice([30, 45, 60])} days of receipt.")
        if ai:
            picked = (p["assumptions"][:2] + rng.sample(p["assumptions"][2:], k=1) if strong
                      else rng.sample(p["assumptions"][2:], k=2))
        else:
            picked = rng.sample(p["assumptions"], k=min(2, len(p["assumptions"])))
        base_assumptions = [f"Assumption: {a}" for a in picked]
        section3 = base_assumptions + extra_assumptions + [f"Out of scope: {p['out_of_scope'][0]}"]
        sla = rng.random() < (0.7 if strong else 0.3)
        eu = name in EU_CLIENTS
        law, courts = (("Ireland", "Ireland") if eu else ("England and Wales", "England and Wales"))
        company_no = f"{rng.randint(1000000, 9999999):08d}"
        msa_date = date(year - 1, rng.randint(1, 12), rng.randint(1, 28))
        defs = [x for x in bp.definitions(strong)
                if (sector == "financial services" or not x.startswith('"Regulator"'))
                and ("subcontracting" in tags or not x.startswith('"Subcontractor"'))]
        reg_clause = (["Regulatory access. The Supplier will give the Client, its auditors and any Regulator access to "
                       "premises, personnel, systems and records relating to the Services as they reasonably require "
                       "to meet regulatory obligations, and will cooperate with any Regulator investigation."]
                      if sector == "financial services" and strong else [])
        audit = ["Audit. The Client may audit the Supplier's records relating to the Charges and the Services once in "
                 "any 12-month period on 30 days' notice" + (", and at any time on reasonable notice if required by a Regulator."
                 if sector == "financial services" else ".")] \
            if (strong or rng.random() < 0.4) else []
        schedule2 = ([f"Schedule 2 payment milestones: {sched}."] if pricing == "fixed_price"
                     else ["Schedule 2 rates: blended day rate as in clause 2.3; expenses at cost with prior approval."])
        credits = (["Service credits. If the Supplier misses the hypercare response times in this clause in any month, "
                    "the Client is entitled to a service credit of 2% of that month's Charges per missed response, capped "
                    "at 10%. Service credits are a price adjustment and not the Client's sole remedy."] if sla and strong else [])
        sla_txt = (["Service levels. During hypercare the Supplier will respond to Severity 1 incidents within one hour and "
                    "resolve or provide a workaround within four hours, and respond to Severity 2 incidents within four hours."]
                   if sla else [])
        cover = [
            f"Statement of Work number SOW-{year}-{100 + i:03d}, version 1.0, issued {start.isoformat()}, prepared by the Supplier engagement lead.",
            f"This SOW is made under the Master Services Agreement between the parties dated {msa_date.isoformat()} (the MSA).",
            f"Supplier: {bp.SUPPLIER[0]}, company number {bp.SUPPLIER[1]}, registered office {bp.SUPPLIER[2]}.",
            f"Client: {name}, company number {company_no}, {'an' if size[0] in 'aeiou' else 'a'} {size} {sector} organisation.",
            bp.FOOTER,
        ]
        present = {"termination", "payment", "acceptance", "change_control", "liability", "ip", "data_protection",
                   "confidentiality", "force_majeure", "governing_law", "anti_bribery"}
        if strong:
            present |= {"indemnity", "insurance", "exit_management", "non_solicitation", "dispute_resolution"}
        if sla:
            present.add("sla")
        if audit:
            present.add("audit")
        def num(sec, items):
            return [f"{sec}.{k + 1} {t}" for k, t in enumerate(items)]
        pages = [
            ("Statement of Work", cover),
            ("1. Definitions", num(1, defs)),
            ("2. Overview and term", num(2, [
                f"Scope. This SOW describes the {product} services the Supplier will provide to {name} (the Client).",
                f"Term. The Services start on {start.isoformat()} and end on {end.isoformat()}, unless terminated earlier under clause 11.",
                f"Contacts. The Client's project sponsor is {contact} ({contact.split()[0].lower()}@{handle}.example, +44 20 7946 {1000 + i:04d}).",
                (f"Charges. The Services are provided at no charge as a Supplier-funded investment with an internal value of {currency} {value:,.0f}; see clause 6." if pricing == "investment"
                 else f"Charges. Total fixed price: {currency} {value:,.0f}, payable in accordance with clause 6 and Schedule 2." if pricing == "fixed_price"
                 else f"Charges. Services are charged on a time and materials basis at a blended day rate of {currency} {rate:,}, with an estimated total of {currency} {value:,.0f}."),
            ])),
            ("3. Deliverables (Schedule 1)", num(3, del_lines)),
            ("4. Assumptions and exclusions", num(4, section3)),
            ("5. Staffing, governance and subcontracting", num(5, [
                f"Key Personnel. The engagement manager and solution architect named in the mobilisation plan are Key Personnel. {(STAFFING_AI if ai else STAFFING).format(dev=rng.choice([2, 3, 4]))} Engagement management is estimated at PM {pm}h.",
                f"Governance. {GOVERNANCE}"] + sector_clauses)),
            ("6. Charges and payment", num(6, [payment] + schedule2 + sla_txt + credits)),
            ("7. Acceptance and non-conformity", num(7, bp.acceptance(strong) + (bp.ai_acceptance(strong) if ai else []))),
            ("8. Change control", num(8, [f"Change control: {CHANGE}"])),
            ("9. Liability, indemnities and insurance", num(9, bp.liability(strong, rng.choice([1, 1.5, 2]), currency, rng.choice([1000000, 2000000, 5000000])))),
            ("10. Data protection, confidentiality and compliance", num(10, bp.compliance(reg_clause + audit)
                                                                          + (bp.ai_governance(strong, eu, sector == "financial services") if ai else []))),
            ("11. Termination and exit", num(11, bp.termination_exit(strong, rng.choice([30, 60, 90]), sector == "financial services")
                                             + (bp.knowledge_transfer(strong) if ai else []))),
            ("12. General", num(12, bp.general(strong, law, courts))),
            ("Signatures", [f"Signed for and on behalf of {bp.SUPPLIER[0]} by its authorised signatory.",
                            f"Signed for and on behalf of {name} by its authorised signatory.", bp.FOOTER]),
        ]
        stem = f"{alias.lower().replace(' ', '_')}_{product.lower().replace(' ', '_')}_{year}"   # files are named by alias
        fname = stem + (".pdf" if i % 3 == 2 else ".docx")
        if fname.endswith(".docx"):
            doc = Document()
            for head, lines in pages:
                doc.add_heading(head, level=1)
                for l in lines:
                    doc.add_paragraph(l)
            doc.save(OUT / "sows" / fname)
        else:
            styles = getSampleStyleSheet()
            story = []
            for head, lines in pages:
                story.append(Paragraph(head, styles["Heading1"]))
                story += [Paragraph(l, styles["Normal"]) for l in lines]
                story.append(PageBreak())
            SimpleDocTemplate(str(OUT / "sows" / fname), pagesize=A4).build(story[:-1])

        register.append({"filename": fname, "sector": sector, "proven": str(health != "red").lower(),
                         "delivery_health": health,
                         "margin_band": {"green": "on", "amber": "below" if not strong else "on", "red": "below"}[health],
                         "csat": {"green": rng.choice([9, 10]), "amber": rng.choice([7, 8]), "red": rng.choice([5, 6])}[health],
                         # Outcomes for pre-sales investments and AI engagements.
                         "converted": (str(strong and health == "green" and i not in NOT_CONVERTED).lower()
                                       if pricing == "investment" else ""),
                         "weeks_to_production": ("" if not ai or health == "red" else
                                                 rng.choice([6, 7, 8, 10]) if strong else rng.choice([14, 18, 22]))})
        if "regulatory" in tags:
            present.add("regulatory")
        if "subcontracting" in tags:
            present.add("subcontracting")
        if ai:
            present |= {"ai_governance", "knowledge_transfer"}
            tags["ai"] = "strong" if strong else "weak"
        facts.append({"filename": fname, "sector": sector, "health": health, "clauses": present, **tags})

        golden.append({"type": "extraction", "filename": fname,
                       "expect": {"pricing_model": pricing, "currency": currency,
                                  "start_date": start.isoformat(), "end_date": end.isoformat()}})
        page_no = next(k for k, t in enumerate(read_pages(OUT / "sows" / fname), 1)
                       if "3.1 The Supplier will deliver:" in " ".join(t.split()))
        golden.append({"type": "retrieval", "query": f"{alias} {product} " + " ".join(dels[0][0].format(n="").split()[:6]),
                       "expect": {"filename": fname, "page": page_no}})
        if i % 4 == 0:
            golden.append({"type": "draft", "brief": {"client_alias": "Client ZZ", "objective": "; ".join(t.split(" (estimated effort")[0].split("deliver: ")[1] for t in del_lines),
                           "products": [product], "size_band": size, "pricing_model": pricing,
                           "exclude_filenames": [fname]},   # hold out the source deal so it can't be copied
                           "reference_hours": sum(effort_total.values())})

    # Clause types whose presence is unambiguous by construction.
    clause_types = ["liability", "sla", "regulatory", "indemnity", "insurance", "non_solicitation", "force_majeure"]
    for ct in clause_types:
        golden.append({"type": "coverage", "clause_type": ct,
                       "expect_missing": [f["filename"] for f in facts if ct not in f["clauses"]]})
    # Precedent cases: the comparison tool must return every deal in the sector that has the clause.
    for ct, sector in (("regulatory", "financial services"), ("acceptance", None),
                       ("payment", "retail"), ("subcontracting", "hospitality"),
                       ("exit_management", None), ("termination", "financial services"),
                       ("audit", "financial services"), ("ai_governance", None), ("knowledge_transfer", None)):
        golden.append({"type": "precedent", "clause_type": ct, "sector": sector,
                       "expect_files": [f["filename"] for f in facts
                                        if ct in f["clauses"] and (sector is None or f["sector"] == sector)]})
    with open(OUT / "alias_map.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["real", "alias"]); w.writerows(aliases)
        w.writerows([(c, "[person]") for c in CONTACTS])
    with open(OUT / "register.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(register[0])); w.writeheader(); w.writerows(register)
    (OUT / "golden.jsonl").write_text("\n".join(json.dumps(g) for g in golden) + "\n")
    with open(OUT / "corpus_facts.json", "w") as f:      # ground truth for checking answers by hand
        json.dump([{**x, "clauses": sorted(x["clauses"])} for x in facts], f, indent=1)
    print(f"{len(CLIENTS)} SOWs, {len(golden)} golden cases -> {OUT}")


if __name__ == "__main__":
    build()
