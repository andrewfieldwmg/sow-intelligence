"""Phase 1: load sanitised SOWs into sow + page. Safe to re-run: files already loaded are skipped."""
import csv
import hashlib

import pdfplumber
from docx import Document

from config import REGISTER_PATH, SOW_PATH, connect
from sanitise import leak_check, load_aliases, sanitise

SUFFIXES = {".pdf", ".docx"}


def read_pages(path) -> list[str]:
    """PDF: one entry per page. Word has no pages, so each Heading 1 section becomes a 'page'."""
    if path.suffix.lower() == ".pdf":
        with pdfplumber.open(path) as pdf:
            return [p.extract_text() or "" for p in pdf.pages]
    doc = Document(path)
    pages, current = [], []
    for para in doc.paragraphs:
        if para.style.name.startswith("Heading 1") and current:
            pages.append("\n".join(current))
            current = []
        if para.text.strip():
            current.append(para.text)
    if current:
        pages.append("\n".join(current))
    # Tables (rate cards, effort grids) become a final page of their own.
    rows = [" | ".join(c.text.strip() for c in row.cells) for t in doc.tables for row in t.rows]
    if rows:
        pages.append("\n".join(rows))
    return pages


def load_register() -> dict:
    """register.csv is maintained by a person: filename,proven,delivery_health,margin_band,csat"""
    if not REGISTER_PATH.exists():
        return {}
    with open(REGISTER_PATH, newline="", encoding="utf-8") as f:
        return {r["filename"]: r for r in csv.DictReader(f)}


def main():
    aliases, register = load_aliases(), load_register()
    files = sorted(p for p in SOW_PATH.rglob("*") if p.suffix.lower() in SUFFIXES)
    conn = connect("writer")
    added = skipped = 0
    with conn, conn.cursor() as cur:
        for path in files:
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            cur.execute("SELECT 1 FROM sow WHERE file_sha256 = %s", (digest,))
            if cur.fetchone():
                skipped += 1
                continue
            pages = [sanitise(t, aliases) for t in read_pages(path)]
            for i, t in enumerate(pages, 1):
                leak_check(t, aliases, f"{path.name} page {i}")
            reg = register.get(path.name, {})
            cur.execute(
                """INSERT INTO sow (filename, file_sha256, proven, delivery_health, margin_band, csat)
                   VALUES (%s, %s, %s, %s, %s, %s) RETURNING id""",
                (path.name, digest, reg.get("proven", "").lower() == "true",
                 reg.get("delivery_health") or None, reg.get("margin_band") or None,
                 reg.get("csat") or None))
            sow_id = cur.fetchone()[0]
            cur.executemany("INSERT INTO page (sow_id, page_no, text) VALUES (%s, %s, %s)",
                            [(sow_id, i, t) for i, t in enumerate(pages, 1)])
            added += 1
    print(f"Ingested {added} new SOWs, skipped {skipped} already loaded")


if __name__ == "__main__":
    main()
