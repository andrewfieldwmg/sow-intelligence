"""Builds three small synthetic SOWs (two Word, one PDF). Fictional content only."""
import sys
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)

def docx(name, sections):
    d = Document()
    for head, lines in sections:
        d.add_heading(head, level=1)
        for l in lines:
            d.add_paragraph(l)
    t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Role"; t.cell(0, 1).text = "Rate"
    t.cell(1, 0).text = "Architect"; t.cell(1, 1).text = "GBP 1,200 per day"
    d.save(out / name)

docx("client_a_aem_2024.docx", [
    ("1. Overview", ["1.1 This scope covers Example Retail Group plc AEM Sites migration, start 2024-04-01 end 2025-03-31.",
                     "1.2 Total fixed price: GBP 480,000 payable by invoice on milestones.",
                     "Contact jane.doe@example.com or +44 20 7946 0000."]),
    ("2. Deliverables", ["2.1 The Supplier will deliver 40 AEM page templates and components (Architect 80h, Developer 600h).",
                         "2.2 The Supplier will deliver a content migration of 12,000 pages to AEM (Developer 400h, PM 120h).",
                         "2.3 Assumption: the Client provides content owners for each market within five days."]),
    ("3. Terms", ["3.1 Either party may terminate for convenience on 30 days written notice.",
                  "3.2 Total liability is capped at the fees paid in the prior twelve months.",
                  "3.3 Weekly steering meetings will review RAID and budget."]),
])
docx("client_b_aem_2023.docx", [
    ("1. Overview", ["1.1 This scope covers Client B AEM Assets rollout, start 2023-01-09 end 2023-09-30.",
                     "1.2 Total fixed price: GBP 210,000 payable by invoice monthly."]),
    ("2. Deliverables", ["2.1 The Supplier will deliver AEM Assets metadata schema and workflows (Architect 40h, Developer 200h, PM 40h).",
                         "2.2 Out of scope: translation of assets and any third-party DAM migration.",
                         "2.3 hallucinate test marker paragraph that is long enough to be a component."]),
    ("3. Terms", ["3.1 Either party may terminate for material breach on 14 days notice."]),
])
c = canvas.Canvas(str(out / "client_c_target_2023.pdf"), pagesize=A4)
for i, l in enumerate(["1.1 This scope covers Client C Target personalisation, start 2023-02-01 end 2023-08-31.",
                       "1.2 The Supplier will deliver 6 personalisation activities in Target (Architect 20h, Developer 90h)."]):
    c.drawString(50, 800 - 20 * i, l)
c.showPage()
c.drawString(50, 800, "1.3 Either party may terminate on 60 days notice. Fees are payable by invoice.")
c.save()
